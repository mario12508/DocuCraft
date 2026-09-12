__all__ = ()

import json
import logging
import re

from docx import Document as DocxDocument

logger = logging.getLogger(__name__)


PLACEHOLDER_RE = re.compile(r"\[([^\[\]\n]{1,60})\]")


KNOWN_FIELDS = {
    "кому": ("addressee", "Адресат"),
    "адресат": ("addressee", "Адресат"),
    "от кого": ("sender", "От кого"),
    "автор": ("sender", "Автор"),
    "отправитель": ("sender", "Отправитель"),
    "составитель": ("sender", "Составитель"),
    "подпись": ("signature", "Подпись"),
    "дата": ("date", "Дата"),
    "номер": ("number", "Номер"),
    "заголовок": ("subject", "Заголовок"),
    "тема": ("subject", "Тема"),
    "текст документа": ("body", "Текст документа"),
    "текст": ("body", "Текст документа"),
    "организация": ("organization", "Организация"),
    "название организации": ("organization", "Организация"),
    "должность": ("sender_position", "Должность"),
    "фио": ("sender_name", "ФИО"),
    "и.о. фамилия": ("sender_name", "ФИО"),
    "и. о. фамилия": ("sender_name", "ФИО"),
}


def extract_raw_placeholders(file_obj):
    """Возвращает список плейсхолдеров из DOCX (без скобок)."""
    file_obj.seek(0)
    docx_doc = DocxDocument(file_obj)

    found = []

    def scan_text(text):
        for m in PLACEHOLDER_RE.finditer(text or ""):
            found.append(m.group(1).strip())

    for p in docx_doc.paragraphs:
        scan_text(p.text)
    for table in docx_doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    scan_text(p.text)
    for section in docx_doc.sections:
        for p in section.header.paragraphs:
            scan_text(p.text)
        for p in section.footer.paragraphs:
            scan_text(p.text)

    seen = set()
    result = []
    for ph in found:
        key = ph.lower()
        if key not in seen:
            seen.add(key)
            result.append(ph)
    return result


def _fallback_map(placeholders):
    """Локальное сопоставление без ИИ."""
    result = []
    for i, ph in enumerate(placeholders, start=1):
        key = ph.strip().lower()
        if key in KNOWN_FIELDS:
            code, label = KNOWN_FIELDS[key]
        else:
            code = f"custom_{i}"
            label = ph.strip()
        result.append(
            {
                "placeholder": f"[{ph}]",
                "code": code,
                "label": label,
            }
        )
    return result


PROMPT = """Ты — эксперт по делопроизводству.
Тебе дан список плейсхолдеров из шаблона служебного документа.
Сопоставь каждый плейсхолдер с кодом и названием реквизита.

Известные коды (используй их, если подходит):
- addressee — Адресат
- sender — От кого / Автор / Составитель / Отправитель
- sender_position — Должность автора
- sender_name — ФИО автора
- signature — Подпись
- date — Дата
- number — Номер
- subject — Заголовок / Тема
- body — Текст документа
- organization — Организация

Если плейсхолдер не подходит ни под один известный код —
предложи новый код (латиница, snake_case) и короткое название.

Верни СТРОГО JSON-объект:
{{
  "items": [
    {{"placeholder": "[Кому]", "code": "addressee", "label": "Адресат"}},
    {{"placeholder": "[Дата]", "code": "date", "label": "Дата"}}
  ]
}}

Плейсхолдеры:
{placeholders}
"""


def _match_via_openai(placeholders, provider):
    """Ветка для OpenAI-совместимых провайдеров (Groq, Gemini, OpenRouter)."""
    from openai import OpenAI

    prompt = PROMPT.format(placeholders=json.dumps(placeholders, ensure_ascii=False))

    client = OpenAI(
        base_url=provider["base_url"],
        api_key=provider["api_key"],
        timeout=provider.get("timeout", 20),
    )
    response = client.chat.completions.create(
        model=provider["model"],
        messages=[
            {"role": "system", "content": "Ты отвечаешь строго JSON-объектом с ключом items."},
            {"role": "user", "content": prompt},
        ],
        temperature=0.1,
        response_format={"type": "json_object"},
    )
    raw = response.choices[0].message.content
    return _parse_items(raw)


def _match_via_gigachat(placeholders, provider):
    """Ветка для GigaChat — своя авторизация через access_token."""
    import uuid

    import requests
    import urllib3

    urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

    auth_key = provider["api_key"]

    # 1) Обмен Authorization key → access_token
    token_resp = requests.post(
        "https://ngw.devices.sberbank.ru:9443/api/v2/oauth",
        headers={
            "Content-Type": "application/x-www-form-urlencoded",
            "Accept": "application/json",
            "RqUID": str(uuid.uuid4()),
            "Authorization": f"Basic {auth_key}",
        },
        data={"scope": "GIGACHAT_API_PERS"},
        verify=False,
        timeout=15,
    )
    token_resp.raise_for_status()
    token = token_resp.json()["access_token"]

    # 2) Запрос на сопоставление плейсхолдеров
    prompt = PROMPT.format(placeholders=json.dumps(placeholders, ensure_ascii=False))
    resp = requests.post(
        "https://api.giga.chat/v1/chat/completions",
        headers={
            "Content-Type": "application/json",
            "Authorization": f"Bearer {token}",
        },
        json={
            "model": provider.get("model", "GigaChat-2"),
            "messages": [
                {"role": "system", "content": "Ты отвечаешь строго JSON-объектом с ключом items."},
                {"role": "user", "content": prompt},
            ],
            "temperature": 0.1,
        },
        verify=False,
        timeout=provider.get("timeout", 30),
    )
    resp.raise_for_status()
    raw = resp.json()["choices"][0]["message"]["content"]
    return _parse_items(raw)


def _parse_items(raw):
    """Общий разбор ответа: dict с items, список или одиночный объект."""
    data = json.loads(raw)
    if isinstance(data, dict):
        items = data.get("items")
        if isinstance(items, list):
            return items
        return [data]
    if isinstance(data, list):
        return data
    return []


def match_placeholders_via_ai(placeholders, provider):
    if not placeholders:
        return []

    if not provider:
        raise ValueError("Не передан провайдер")

    name = provider.get("name", "")

    if name == "gigachat":
        return _match_via_gigachat(placeholders, provider)

    if provider.get("base_url"):
        return _match_via_openai(placeholders, provider)

    raise ValueError(f"Провайдер {name} не поддерживается для сопоставления плейсхолдеров")


def parse_template(file_obj, provider=None):
    """
    Возвращает (placeholders, error).
    placeholders — список словарей [{"placeholder", "code", "label"}].
    """
    raw = extract_raw_placeholders(file_obj)
    if not raw:
        return [], "В файле не найдено ни одного плейсхолдера вида [Кому]."

    if provider:
        try:
            return match_placeholders_via_ai(raw, provider), ""
        except Exception as exc:
            logger.exception("Ошибка ИИ-сопоставления плейсхолдеров")
            return (
                _fallback_map(raw),
                f"ИИ недоступен, использован базовый словарь: {exc}",
            )

    return _fallback_map(raw), ""
