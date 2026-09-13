__all__ = ()

import logging
import re
from io import BytesIO

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

logger = logging.getLogger(__name__)

ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

BODY_CODES = {"body", "processed_text", "document_text", "text"}

NUMBER_SUFFIX = {
    "sluzhebnaya_zapiska": "СЗ",
    "dokladnaya_zapiska": "ДЗ",
    "pismo": "П",
    "informacionnaya_spravka": None,
}

DOCUMENT_TITLE = {
    "sluzhebnaya_zapiska": "СЛУЖЕБНАЯ ЗАПИСКА",
    "dokladnaya_zapiska": "ДОКЛАДНАЯ ЗАПИСКА",
    "informacionnaya_spravka": "ИНФОРМАЦИОННАЯ СПРАВКА",
    "pismo": "ПИСЬМО",
}


def _normalize_text(text):
    """Убирает лишние переносы и пробелы."""
    if not text:
        return ""
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    text = re.sub(r"[ \t]{2,}", " ", text)
    return text.strip()


def _collect_replacements(document):
    extracted = document.extracted_fields or {}
    tmpl = document.template

    mapping = {}
    for p in tmpl.placeholders or []:
        ph = p.get("placeholder")
        code = p.get("code")
        if not ph or not code:
            continue
        if code in BODY_CODES:
            mapping[ph] = document.processed_text or ""
        else:
            mapping[ph] = str(extracted.get(code) or "")
    return mapping


def _replace_in_paragraph(paragraph, replacements):
    if not paragraph.runs:
        return
    full_text = "".join(run.text for run in paragraph.runs)
    if not any(ph in full_text for ph in replacements):
        return
    new_text = full_text
    for ph, value in replacements.items():
        if ph not in new_text:
            continue
        new_text = new_text.replace(ph, value)
    first_run = paragraph.runs[0]
    first_run.text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""


def _replace_in_table(table, replacements):
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                _replace_in_paragraph(p, replacements)
            for nested in cell.tables:
                _replace_in_table(nested, replacements)


def _render_from_file(document):
    tmpl = document.template
    if not tmpl.docx_template:
        return None
    try:
        tmpl.docx_template.open("rb")
        docx_doc = DocxDocument(tmpl.docx_template)
    finally:
        try:
            tmpl.docx_template.close()
        except Exception:
            pass

    replacements = _collect_replacements(document)
    for p in docx_doc.paragraphs:
        _replace_in_paragraph(p, replacements)
    for table in docx_doc.tables:
        _replace_in_table(table, replacements)
    for section in docx_doc.sections:
        for p in section.header.paragraphs:
            _replace_in_paragraph(p, replacements)
        for p in section.footer.paragraphs:
            _replace_in_paragraph(p, replacements)

    buffer = BytesIO()
    docx_doc.save(buffer)
    buffer.seek(0)
    return buffer




def _apply_page_rules(docx_doc, rules):
    page = rules.get("page", {})
    margins = page.get("margins", {})
    section = docx_doc.sections[0]
    section.top_margin = Cm(margins.get("top", 2))
    section.bottom_margin = Cm(margins.get("bottom", 2))
    section.left_margin = Cm(margins.get("left", 3))
    section.right_margin = Cm(margins.get("right", 1.5))


def _style_run(run, font_name, font_size, bold=False):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for attr in ("w:eastAsia", "w:cs", "w:hAnsi", "w:ascii"):
        rfonts.set(qn(attr), font_name)


def _add_paragraph(
    docx_doc, text, rules, font_name, font_size, align=None, first_line_indent=None, bold=False
):
    """Добавляет абзац. Если text пустой — ничего не делает."""
    if not text:
        return None
    p = docx_doc.add_paragraph(text)
    for run in p.runs:
        _style_run(run, font_name, font_size, bold=bold)
    if align:
        p.alignment = ALIGN_MAP.get(align, WD_ALIGN_PARAGRAPH.JUSTIFY)
    pf = p.paragraph_format
    pf.line_spacing = rules.get("line_spacing", 1.5)
    pf.keep_together = False
    pf.keep_with_next = False
    if first_line_indent is not None:
        pf.first_line_indent = Cm(first_line_indent)
    return p


def _apply_header_footer(docx_doc, rules, document):
    section = docx_doc.sections[0]

    org_name = rules.get("organization")
    if org_name:
        p = section.header.paragraphs[0]
        p.text = org_name
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.size = Pt(rules.get("header_font_size", 11))

    footer_template = rules.get("footer", "")
    if footer_template and document is not None:
        text = footer_template.format(
            document_type=document.document_type.name,
            date=(document.document_date.strftime("%d.%m.%Y") if document.document_date else ""),
        )
        p = section.footer.paragraphs[0]
        p.text = text
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.size = Pt(rules.get("footer_font_size", 10))


def _get_number(document):
    """Номер из extracted_fields или сгенерированный."""
    extracted = document.extracted_fields or {}
    if extracted.get("number"):
        return extracted["number"]
    suffix = NUMBER_SUFFIX.get(document.document_type.code)
    if not suffix:
        return None
    from apps.documents.models import Document as DocModel

    count = DocModel.objects.filter(document_type=document.document_type).count()
    return f"{count + 1:02d}-{suffix}"


def _add_empty_paragraph(docx_doc, font_name, font_size, rules):
    """Пустой абзац с явно нулевыми отступами."""
    p = docx_doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = rules.get("line_spacing", 1.5)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    run = p.add_run("")
    _style_run(run, font_name, font_size)
    return p


def _setup_document_styles(docx_doc, font_name, font_size, line_spacing):
    """
    Настраивает стиль Normal для всего документа:
    - шрифт, кегль
    - межстрочный интервал
    - нулевые отступы до/после абзаца
    """
    style = docx_doc.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(font_size)

    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for attr in ("w:eastAsia", "w:cs", "w:hAnsi", "w:ascii"):
        rfonts.set(qn(attr), font_name)

    pf = style.paragraph_format
    pf.line_spacing = line_spacing
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.widow_control = False


def _render_programmatic(document):
    """Программная сборка с ветвлением по типу документа."""
    rules = document.template.rules or {}
    font = rules.get("font", {})
    font_name = font.get("name", "Times New Roman")
    font_size = font.get("size", 14)
    extracted = document.extracted_fields or {}
    doc_type_code = document.document_type.code
    number = _get_number(document)

    is_spravka = doc_type_code == "informacionnaya_spravka"
    is_pismo = doc_type_code == "pismo"

    docx_doc = DocxDocument()
    _apply_page_rules(docx_doc, rules)
    _setup_document_styles(
        docx_doc,
        font_name,
        font_size,
        rules.get("line_spacing", 1.5),
    )
    _apply_header_footer(docx_doc, rules, document)

    if not is_spravka:
        addressee_parts = [
            extracted.get("addressee_position"),
            extracted.get("addressee_org"),
            extracted.get("addressee_name"),
        ]
        addressee_parts = [p for p in addressee_parts if p]

        if not addressee_parts and extracted.get("addressee"):
            addressee_parts = [s.strip() for s in extracted["addressee"].split("\n") if s.strip()]

        for line in addressee_parts:
            _add_paragraph(
                docx_doc,
                line,
                rules,
                font_name,
                font_size,
                align="right",
                first_line_indent=0,
            )
    date = extracted.get("date", "")
    parts = []
    if date:
        parts.append(f"Дата: {date}")
    if number and not is_spravka:
        parts.append(f"Номер: {number}")
    if parts:
        _add_paragraph(
            docx_doc,
            "   ".join(parts),
            rules,
            font_name,
            font_size,
            align="left",
            first_line_indent=0,
        )

    title = DOCUMENT_TITLE.get(doc_type_code)
    if title:
        _add_paragraph(
            docx_doc,
            title,
            rules,
            font_name,
            font_size,
            align="center",
            first_line_indent=0,
            bold=True,
        )

    subject = extracted.get("subject", "")
    if subject:
        _add_paragraph(
            docx_doc,
            subject,
            rules,
            font_name,
            font_size,
            align="center",
            first_line_indent=0,
            bold=True,
        )
    body = _normalize_text(document.processed_text or "")
    for block in body.split("\n\n"):
        text = block.strip()
        if not text:
            continue
        _add_paragraph(
            docx_doc,
            text,
            rules,
            font_name,
            font_size,
            align="justify",
            first_line_indent=1.25,
        )

    sender_position = extracted.get("sender_position")
    sender_name = extracted.get("sender_name")

    if not sender_position and not sender_name:
        sender_raw = extracted.get("signature") or extracted.get("sender")
        if sender_raw:
            lines = [s.strip() for s in sender_raw.split("\n") if s.strip()]
            if len(lines) >= 2:
                sender_position, sender_name = lines[0], lines[-1]
            elif lines:
                sender_position = lines[0]

    if sender_position or sender_name:
        _add_empty_paragraph(docx_doc, font_name, font_size, rules)

        align = "center" if is_pismo else "left"

        if sender_position:
            _add_paragraph(
                docx_doc,
                sender_position,
                rules,
                font_name,
                font_size,
                align=align,
                first_line_indent=0,
            )
            _add_empty_paragraph(docx_doc, font_name, font_size, rules)
            _add_paragraph(
                docx_doc,
                "___________________________",
                rules,
                font_name,
                font_size,
                align=align,
                first_line_indent=0,
            )
            if sender_name:
                _add_paragraph(
                    docx_doc,
                    sender_name,
                    rules,
                    font_name,
                    font_size,
                    align=align,
                    first_line_indent=0,
                )
        elif sender_name:
            _add_paragraph(
                docx_doc,
                sender_name,
                rules,
                font_name,
                font_size,
                align=align,
                first_line_indent=0,
            )

    buffer = BytesIO()
    docx_doc.save(buffer)
    buffer.seek(0)
    return buffer


def generate_docx(document):
    """Возвращает BytesIO с готовым DOCX."""
    tmpl = document.template

    if tmpl.kind == "user" and tmpl.docx_template:
        from_file = _render_from_file(document)
        if from_file is not None:
            return from_file

    return _render_programmatic(document)
