"""Тесты локального сопоставления плейсхолдеров (без ИИ)."""

from apps.documents.services.template_parser import (
    _fallback_map,
    extract_raw_placeholders,
)


class TestFallbackMap:
    """Локальный словарь должен распознавать стандартные плейсхолдеры."""

    def test_known_placeholders(self):
        result = _fallback_map(["Кому", "Дата", "Номер"])
        codes = {item["code"] for item in result}
        assert codes == {"addressee", "date", "number"}

    def test_case_insensitive(self):
        result = _fallback_map(["КОМУ", "Дата", "НОМЕР"])
        codes = {item["code"] for item in result}
        assert codes == {"addressee", "date", "number"}

    def test_unknown_becomes_custom(self):
        result = _fallback_map(["Кто-то-странный"])
        assert result[0]["code"].startswith("custom_")

    def test_empty_input(self):
        assert _fallback_map([]) == []


class TestExtractPlaceholders:
    """Извлечение плейсхолдеров из DOCX — на базе реального шаблона."""

    def test_no_placeholders(self, tmp_path):
        from docx import Document as Docx

        path = tmp_path / "empty.docx"
        doc = Docx()
        doc.add_paragraph("Просто текст без плейсхолдеров.")
        doc.save(str(path))

        with open(path, "rb") as f:
            result = extract_raw_placeholders(f)

        assert result == []

    def test_finds_placeholders(self, tmp_path):
        from docx import Document as Docx

        path = tmp_path / "tpl.docx"
        doc = Docx()
        doc.add_paragraph("Кому: [Кому]")
        doc.add_paragraph("Дата: [Дата]")
        doc.add_paragraph("[Текст документа]")
        doc.save(str(path))

        with open(path, "rb") as f:
            result = extract_raw_placeholders(f)

        assert set(result) == {"Кому", "Дата", "Текст документа"}

    def test_deduplicates(self, tmp_path):
        from docx import Document as Docx

        path = tmp_path / "tpl.docx"
        doc = Docx()
        doc.add_paragraph("[Дата]")
        doc.add_paragraph("[Дата]")
        doc.add_paragraph("[дата]")
        doc.save(str(path))

        with open(path, "rb") as f:
            result = extract_raw_placeholders(f)

        assert len(result) == 1
