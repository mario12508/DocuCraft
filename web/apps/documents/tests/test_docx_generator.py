"""Тесты программной генерации DOCX."""

from datetime import date

import pytest
from docx import Document as Docx

from apps.documents.docx_generator import generate_docx
from apps.documents.models import Document, DocumentType, Template


@pytest.fixture
def sample_document(db):
    dt = DocumentType.objects.create(
        code="test_type", name="Тестовый тип", order=1,
    )
    tpl = Template.objects.create(
        code="test_tpl", kind="system", name="Тестовый шаблон",
        rules={
            "font": {"name": "Times New Roman", "size": 14},
            "line_spacing": 1.5,
            "paragraph": {"first_line_indent": 1.25, "alignment": "justify"},
            "organization": "ООО «Тест»",
        },
        placeholders=[
            {"placeholder": "[Кому]", "code": "addressee_position",
             "label": "Адресат"},
            {"placeholder": "[Текст документа]", "code": "body",
             "label": "Текст"},
        ],
    )
    return Document.objects.create(
        document_type=dt,
        template=tpl,
        source_text="оригинал",
        processed_text="Обработанный текст документа.",
        extracted_fields={
            "addressee_position": "Директору",
            "addressee_org": "ООО «Ромашка»",
            "addressee_name": "Иванову И.И.",
            "sender_position": "Начальник отдела",
            "sender_name": "Петров П.П.",
            "date": "12.03.2025",
        },
        missing_fields=[],
        status="ready",
        document_date=date(2025, 3, 12),
    )


@pytest.mark.django_db
class TestDocxGenerator:

    def test_output_is_valid_zip(self, sample_document):
        buf = generate_docx(sample_document)
        data = buf.getvalue()
        assert data[:2] == b"PK"
        assert len(data) > 1000

    def test_contains_processed_text(self, sample_document):
        buf = generate_docx(sample_document)
        doc = Docx(buf)
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Обработанный текст" in text

    def test_contains_addressee_parts(self, sample_document):
        buf = generate_docx(sample_document)
        doc = Docx(buf)
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Директору" in text
        assert "ООО «Ромашка»" in text
        assert "Иванову И.И." in text

    def test_no_unresolved_placeholders(self, sample_document):
        buf = generate_docx(sample_document)
        doc = Docx(buf)
        text = "\n".join(p.text for p in doc.paragraphs)
        for ph in ("[Кому]", "[Текст документа]", "[Дата]"):
            assert ph not in text

    def test_contains_signature(self, sample_document):
        buf = generate_docx(sample_document)
        doc = Docx(buf)
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "Начальник отдела" in text
        assert "Петров П.П." in text
